"""Turn PDF, DOCX, HTML, Markdown and text files into a list of Sections.

Each parser keeps the structure we need for good citations:
page numbers for PDFs and headings for everything else.
"""
import hashlib
import re
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.ingestion.models import Document, Section


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"-\n(\w)", r"\1", text)      # re-join words split across lines
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(path: Path) -> list[Section]:
    reader = PdfReader(str(path))
    sections = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            sections.append(Section(text=text, page=page_no))
    return sections


def parse_docx(path: Path) -> list[Section]:
    doc = DocxDocument(str(path))
    sections: list[Section] = []
    heading, buffer = None, []

    def flush():
        if buffer:
            sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
            buffer.clear()

    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        if para.style.name.lower().startswith("heading"):
            flush()
            heading = para.text.strip()
        else:
            buffer.append(para.text)
    flush()
    return sections


def parse_html(path: Path) -> list[Section]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    sections: list[Section] = []
    heading, buffer = None, []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "td"]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name in {"h1", "h2", "h3", "h4"}:
            if buffer:
                sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
                buffer = []
            heading = text
        else:
            buffer.append(text)
    if buffer:
        sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
    return sections


def parse_markdown(path: Path) -> list[Section]:
    sections: list[Section] = []
    heading, buffer = None, []
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        match = re.match(r"^#{1,6}\s+(.*)", line)
        if match:
            if "".join(buffer).strip():
                sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
            heading, buffer = match.group(1).strip(), []
        else:
            buffer.append(line)
    if "".join(buffer).strip():
        sections.append(Section(text=_clean("\n".join(buffer)), heading=heading))
    return sections


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".html": parse_html,
    ".htm": parse_html,
    ".md": parse_markdown,
    ".txt": parse_markdown,
}


def load_document(path: str | Path, allowed_roles: list[str] | None = None) -> Document:
    path = Path(path)
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    # Stable ID from the file contents: re-uploading the same file won't duplicate it.
    doc_id = hashlib.sha256(path.read_bytes()).hexdigest()[:16]
    return Document(
        doc_id=doc_id,
        source=path.name,
        sections=parser(path),
        allowed_roles=allowed_roles or ["all"],
    )
