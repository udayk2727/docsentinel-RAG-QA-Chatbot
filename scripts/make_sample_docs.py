"""Generate demo documents for the fictional company 'Northwind Labs'.

Different files belong to different departments so we can demo access control later.
"""
import json
from pathlib import Path

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).resolve().parent.parent / "data" / "sample_docs"
OUT.mkdir(parents=True, exist_ok=True)

(OUT / "employee_handbook.md").write_text("""# Northwind Labs Employee Handbook

## Leave Policy
Full-time employees receive 24 days of paid annual leave per calendar year. Up to 5 unused days can be carried over to the next year. Leave requests must be submitted in the HR portal at least 10 working days in advance.

## Remote Work
Employees may work remotely up to 3 days per week with manager approval. Fully remote arrangements require VP approval and are reviewed every 6 months.

## Expense Reimbursement
Business expenses must be submitted within 30 days with receipts. Meals while travelling are reimbursed up to $60 per day. Expenses above $500 need director approval.
""")

doc = DocxDocument()
doc.add_heading("Compensation Bands 2026 (Confidential)", level=1)
doc.add_heading("Engineering", level=2)
doc.add_paragraph("Software Engineer II base salary band is $118,000 to $142,000. Senior Software Engineer band is $150,000 to $185,000.")
doc.add_heading("Bonus Policy", level=2)
doc.add_paragraph("The annual target bonus is 10% of base salary for individual contributors and 15% for managers. Bonuses are paid in March.")
doc.save(OUT / "compensation_bands.docx")

(OUT / "incident_runbook.html").write_text("""<html><body>
<h1>Incident Response Runbook</h1>
<h2>Severity Levels</h2>
<p>SEV1 means customer-facing outage affecting more than 20% of users. SEV2 means degraded performance or a partial outage. SEV3 means internal tooling issues.</p>
<h2>On-call Escalation</h2>
<p>For a SEV1, page the on-call engineer immediately via PagerDuty and open a bridge call within 15 minutes. The incident commander must post updates in #incidents every 30 minutes.</p>
<h2>Postmortems</h2>
<p>A blameless postmortem is required for every SEV1 and SEV2 within 5 business days.</p>
</body></html>""")

c = canvas.Canvas(str(OUT / "security_policy.pdf"), pagesize=A4)
pages = [
    ["Information Security Policy", "", "1. Passwords",
     "Passwords must be at least 14 characters long.",
     "Multi-factor authentication is mandatory for all company accounts."],
    ["2. Data Classification",
     "Data is classified as Public, Internal, Confidential, or Restricted.",
     "Restricted data must never be stored on personal devices.",
     "", "3. Reporting",
     "Report suspected phishing to security@northwind.example within 1 hour."],
]
for lines in pages:
    y = 800
    for line in lines:
        c.drawString(60, y, line)
        y -= 20
    c.showPage()
c.save()

acl = {
    "employee_handbook.md": ["all"],
    "security_policy.pdf": ["all"],
    "incident_runbook.html": ["engineering", "admin"],
    "compensation_bands.docx": ["hr", "admin"],
}
(OUT / "acl.json").write_text(json.dumps(acl, indent=2))
print("Sample docs written to", OUT)
